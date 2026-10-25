"""Exports every table of workorders.db into one Excel file (one sheet per table)."""
import os, sqlite3, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font

here = os.path.dirname(os.path.abspath(__file__))
db_path = os.path.join(here, "workorders.db")
if not os.path.exists(db_path):
    print("workorders.db not found - start the app once with run.bat first.")
    sys.exit(1)

HIDE = {"password_hash", "file_data"}   # password hashes and file contents are left out
con = sqlite3.connect("file:" + db_path + "?mode=ro", uri=True)
wb = Workbook(); wb.remove(wb.active)
tables = [r[0] for r in con.execute("select name from sqlite_master where type='table' and name not like 'sqlite_%' order by name")]
for t in tables:
    cur = con.execute(f'select * from "{t}"')
    cols = [d[0] for d in cur.description]
    keep = [i for i, c in enumerate(cols) if c not in HIDE]
    ws = wb.create_sheet(t[:31])
    ws.append([cols[i] for i in keep])
    n = 0
    for row in cur:
        out = []
        for i in keep:
            v = row[i]
            if isinstance(v, (bytes, bytearray)): v = "(binary)"
            elif isinstance(v, str) and len(v) > 30000: v = v[:30000] + "...(cut)"
            out.append(v)
        ws.append(out); n += 1
    ws.freeze_panes = "A2"
    for c in ws[1]: c.font = Font(bold=True)
    print(f"{t}: {n} rows")
out = os.path.join(here, "workorders_tables.xlsx")
wb.save(out)
print("\nSaved:", out)
